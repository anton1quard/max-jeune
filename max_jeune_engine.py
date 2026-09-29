"""
===============================================================================
 MAX JEUNE - Moteur de recherche (V2) — module moteur, sans interface
===============================================================================

Ce module ne dépend que de la bibliothèque standard (openpyxl est optionnel,
utilisé seulement pour lire l'Excel des axes, réservé à la future Partie D).
Il est importé par l'interface Tkinter (max_jeune.py), par les tests
(tests/test_engine.py) et servira de référence pour la future version web.

PRINCIPE CLÉ DE LA V2
    Le CSV TGVMax contient, pour chaque train et chaque jour, plusieurs
    lignes (sous-trajets) avec le statut OUI **ou NON**. En gardant toutes
    les lignes, on reconstitue pour chaque circulation (date, train) :
      - les gares réellement desservies,
      - l'heure réelle de passage à chaque gare,
      - l'ordre de desserte (graphe "gare A -> gare B" issu des lignes).
    La montée en cours d'arrêt (Partie B) et les correspondances (Partie C)
    s'appuient sur ces données réelles : plus d'axes saisis à la main, plus
    d'horaires estimés, et les deux sens de circulation sont couverts.

GARDE-FOU "TRAINS QUI SE SÉPARENT"
    Un même numéro de train peut desservir deux branches (ex. une rame vers
    Toulon, l'autre vers Nice). On ne trie donc pas bêtement les gares par
    heure : une gare S n'est considérée sur le trajet du billet O -> D que si
    S est atteignable depuis O **et** D atteignable depuis S dans le graphe
    des lignes du train. Deux gares de branches différentes ne sont jamais
    reliées par une ligne du CSV, donc jamais mélangées.

Structure :
    1. IMPORTS
    2. CONFIGURATION (chemins, hubs, marges, alias)
    3. NORMALISATION DES NOMS DE GARES
    4. AXES EXCEL (conservés pour la Partie D, non utilisés en A/B/C)
    5. CHARGEMENT DU CSV -> Dataset (billets OUI + circulations reconstruites)
    6. RÉSOLUTION DES SAISIES UTILISATEUR (gares / hubs)
    7. TRONÇONS RÉSERVABLES ("legs")
    8. RECHERCHE : Parties A, B, C + filtre de dominance
    9. PARTIE D : MAX + TER (GTFS SNCF)
    10. POINT D'ENTRÉE DE LA RECHERCHE
    11. BALAYAGE MULTI-JOURS
"""

# =============================================================================
# 1. IMPORTS
# =============================================================================
from __future__ import annotations

import csv
import io
import os
import zipfile
from bisect import bisect_left, bisect_right
import re
import time
import unicodedata
import urllib.request
from collections import defaultdict
from dataclasses import dataclass, field, replace
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Dict, Iterable, List, NamedTuple, Optional, Sequence, Set, Tuple


# =============================================================================
# 2. CONFIGURATION
# =============================================================================

BASE_DIR = Path(__file__).resolve().parent
AXES_XLSX_PATH = BASE_DIR / "axes_max_jeune.xlsx"
TGVMAX_CACHE_PATH = BASE_DIR / "tgvmax_cache.csv"
TGVMAX_URL = (
    "https://ressources.data.sncf.com/api/explore/v2.1/catalog/datasets/"
    "tgvmax/exports/csv?lang=fr&timezone=Europe%2FBerlin"
    "&use_labels=true&delimiter=%3B"
)
CACHE_MAX_AGE_HOURS = 12

# --- Hubs ---------------------------------------------------------------------
HUBS: Dict[str, List[str]] = {
    "PARIS": [
        "PARIS GARE DE LYON", "PARIS MONTPARNASSE", "PARIS NORD", "PARIS EST",
        "PARIS AUSTERLITZ", "MARNE LA VALLEE CHESSY", "AEROPORT CDG 2 TGV",
        "MASSY TGV", "MASSY PALAISEAU", "VERSAILLES CHANTIERS",
    ],
    "LYON": ["LYON PART DIEU", "LYON PERRACHE", "LYON SAINT-EXUPERY TGV"],
    "LYON ETENDU": [
        "LYON PART DIEU", "LYON PERRACHE", "LYON SAINT-EXUPERY TGV",
        "LE CREUSOT TGV", "MACON LOCHE TGV", "BOURG EN BRESSE", "VIENNE",
        "CHALON SUR SAONE", "DIJON VILLE", "AMBERIEU", "GRENOBLE",
        "CHAMBERY CHALLES LES EAUX", "VALENCE TGV RHONE-ALPES SUD",
        "VALENCE VILLE", "ST ETIENNE CHATEAUCREUX",
    ],
    # Marseille : St Charles seul par défaut ; Aix TGV avec l'option "étendu"
    # (navette bus Aix TGV <-> St Charles ~ 30 min).
    "MARSEILLE": ["MARSEILLE ST CHARLES"],
    "MARSEILLE ETENDU": ["MARSEILLE ST CHARLES", "AIX EN PROVENCE TGV"],
}
HUB_EXTENDED_MAP: Dict[str, str] = {
    "LYON": "LYON ETENDU",
    "MARSEILLE": "MARSEILLE ETENDU",
}

# Hubs où l'on accepte de changer de gare entre deux segments (Partie C).
CONNECTION_HUBS: List[str] = ["PARIS", "LYON"]
HUB_SIBLINGS: Dict[str, Set[str]] = {}
STATION_TO_HUB: Dict[str, str] = {}
for _h in CONNECTION_HUBS:
    for _s in HUBS[_h]:
        HUB_SIBLINGS[_s] = set(HUBS[_h])
        STATION_TO_HUB[_s] = _h

# --- Marges de correspondance (Partie C), en minutes -------------------------
MARGIN_SAME_TRAIN_MAX = 60      # même train, même gare : 0 à 60 min
MARGIN_SAME_STATION_MIN = 10    # trains différents, même gare
MARGIN_HUB_MIN = 60             # changement de gare dans un hub
MARGIN_CONNECTION_MAX = 240     # plafond de correspondance
TWO_SEG_TOTAL_DURATION_MAX_MIN = 600   # durée totale max d'un trajet 2 segments

SCAN_DEFAULT_DAYS = 30

# --- Alias saisie utilisateur -> nom canonique --------------------------------
MANUAL_ALIASES: Dict[str, str] = {
    "PARIS LYON": "PARIS GARE DE LYON",
    "PARIS GDL": "PARIS GARE DE LYON",
    "CDG 2 TGV": "AEROPORT CDG 2 TGV",
    "CHARLES DE GAULLE": "AEROPORT CDG 2 TGV",
    "ROISSY CDG": "AEROPORT CDG 2 TGV",
    "MARNE LA VALLEE": "MARNE LA VALLEE CHESSY",
    "LYON ST EXUPERY": "LYON SAINT-EXUPERY TGV",
    "LYON SAINT EXUPERY": "LYON SAINT-EXUPERY TGV",
    "AIX TGV": "AIX EN PROVENCE TGV",
    "AVIGNON": "AVIGNON TGV",
    "VALENCE TGV": "VALENCE TGV RHONE-ALPES SUD",
    "MONTPELLIER": "MONTPELLIER SUD DE FRANCE",
    "BORDEAUX": "BORDEAUX ST JEAN",
    "LILLE": "LILLE EUROPE",
    "NIMES": "NIMES CENTRE",
}

# --- Alias noms CSV -> nom canonique -------------------------------------------
_HALL_SUFFIXES = [" HALL 1 ET 2", " HALL 1 2", " HALL 12", " HALL 3", " HALL 2", " HALL 1"]

STATION_CSV_ALIASES: Dict[str, str] = {
    "AEROPORT CHARLES DE GAULLE 2 TGV": "AEROPORT CDG 2 TGV",
    "AEROPORT CHARLES DE GAULLE TGV": "AEROPORT CDG 2 TGV",
    "AEROPORT ROISSY CDG 2 TGV": "AEROPORT CDG 2 TGV",
    "AEROPORT ROISSY CHARLES DE GAULLE 2 TGV": "AEROPORT CDG 2 TGV",
    "VALENCE TGV": "VALENCE TGV RHONE-ALPES SUD",
    "VALENCE TGV AUVERGNE RHONE ALPES": "VALENCE TGV RHONE-ALPES SUD",
    "VALENCE TGV AUVERGNE RHONE-ALPES": "VALENCE TGV RHONE-ALPES SUD",
    "LE CREUSOT MONTCEAU MONTCHANIN": "LE CREUSOT TGV",
    "LE CREUSOT MONTCEAU MONTCHANIN TGV": "LE CREUSOT TGV",
    "MARSEILLE ST-CHARLES": "MARSEILLE ST CHARLES",
    "LYON ST EXUPERY TGV": "LYON SAINT-EXUPERY TGV",
    "LYON ST-EXUPERY TGV": "LYON SAINT-EXUPERY TGV",
    "LYON SAINT EXUPERY TGV": "LYON SAINT-EXUPERY TGV",
    "SAINT ETIENNE CHATEAUCREUX": "ST ETIENNE CHATEAUCREUX",
    "MONTPELLIER SAINT ROCH": "MONTPELLIER ST-ROCH",
    "MONTPELLIER ST ROCH": "MONTPELLIER ST-ROCH",
    "BORDEAUX SAINT JEAN": "BORDEAUX ST JEAN",
    "ANGERS SAINT LAUD": "ANGERS ST LAUD",
    "SAINT PIERRE DES CORPS": "ST PIERRE DES CORPS",
    # Libellés utilisés dans l'Excel des axes
    "NIMES": "NIMES CENTRE",
}

# Code IATA -> gare. Indispensable pour les libellés "VILLE (intramuros)"
# qui regroupent plusieurs gares sous un même nom dans le CSV.
IATA_TO_CANONICAL: Dict[str, str] = {
    "FRPLY": "PARIS GARE DE LYON",
    "FRPMO": "PARIS MONTPARNASSE",
    "FRPNO": "PARIS NORD",
    "FRPST": "PARIS EST",
    "FRPAZ": "PARIS AUSTERLITZ",
    "FRPBE": "PARIS AUSTERLITZ",     # Paris Bercy assimilée à Austerlitz
    "FRLPD": "LYON PART DIEU",
    "FRLPE": "LYON PERRACHE",
    "FRJDQ": "LYON SAINT-EXUPERY TGV",
    # Lille : déduit des dessertes (FRLLE = Bruxelles/Lyon/Marseille,
    # FRADJ = principalement Paris Nord)
    "FRLLE": "LILLE EUROPE",
    "FRADJ": "LILLE FLANDRES",
}


# =============================================================================
# 3. NORMALISATION DES NOMS DE GARES
# =============================================================================

def strip_accents(text: str) -> str:
    nfkd = unicodedata.normalize("NFKD", text)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def normalize(text: str) -> str:
    """MAJUSCULES, sans accents, ponctuation réduite, espaces simples."""
    if text is None:
        return ""
    t = strip_accents(str(text)).upper()
    t = re.sub(r"[^A-Z0-9\-\s]", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def resolve_alias(user_text: str) -> str:
    n = normalize(user_text)
    return MANUAL_ALIASES.get(n, n)


@lru_cache(maxsize=None)
def canonicalize_station(raw: str) -> str:
    """
    Nom brut (CSV ou Excel) -> nom canonique.
    Les parenthèses sont retirées AVANT normalisation (sinon normalize les
    transforme en espaces et "(intramuros)" restait collé au nom : bug V1).
    """
    if not raw:
        return ""
    n = normalize(re.sub(r"\([^)]*\)", " ", str(raw)))
    if n in STATION_CSV_ALIASES:
        return STATION_CSV_ALIASES[n]
    for suf in _HALL_SUFFIXES:
        if n.endswith(suf):
            n = n[: -len(suf)].strip()
            break
    return STATION_CSV_ALIASES.get(n, n)


@lru_cache(maxsize=None)
def resolve_csv_station(name: str, iata: str) -> str:
    """Le code IATA est prioritaire (désambiguïse les "intramuros")."""
    code = (iata or "").strip().upper()
    if code in IATA_TO_CANONICAL:
        return IATA_TO_CANONICAL[code]
    return canonicalize_station(name)


# =============================================================================
# 4. AXES EXCEL (conservés pour la future Partie D — TER)
# =============================================================================

def _split_stops_cell(cell: str) -> List[str]:
    parts = re.split(r"\s*(?:→|->|›|»)\s*", cell or "")
    return [p.strip() for p in parts if p.strip()]


def _expand_or_variants(stops: Sequence[str]) -> List[List[str]]:
    """'X OU Y' dans un arrêt -> une variante de pattern par alternative."""
    variants: List[List[str]] = [[]]
    for stop in stops:
        alts = [a.strip() for a in re.split(r"\s+OU\s+", stop, flags=re.IGNORECASE)]
        variants = [v + [a] for v in variants for a in alts]
    return variants


def load_patterns_from_excel(path: Path) -> Dict[Tuple[str, str], List[List[str]]]:
    """(origine, destination) -> [patterns]. Nécessite openpyxl."""
    import openpyxl  # import local : dépendance optionnelle

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Axes"] if "Axes" in wb.sheetnames else wb.active
    patterns: Dict[Tuple[str, str], List[List[str]]] = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not row or len(row) < 4 or not all(row[1:4]):
            continue
        o, d = canonicalize_station(row[1]), canonicalize_station(row[2])
        for variant in _expand_or_variants(_split_stops_cell(str(row[3]))):
            pat = [canonicalize_station(s) for s in variant]
            if pat[0] != o:
                pat.insert(0, o)
            if pat[-1] != d:
                pat.append(d)
            patterns.setdefault((o, d), []).append(pat)
    return patterns


# =============================================================================
# 5. CHARGEMENT DU CSV -> Dataset
# =============================================================================

@dataclass(frozen=True)
class MaxTrip:
    """Un billet MAX réservable = une ligne OUI du CSV."""
    date_str: str
    train_no: str
    axe: str
    origine: str
    destination: str
    heure_depart: str
    heure_arrivee: str
    datetime_depart: datetime
    datetime_arrivee: datetime


class TrainRun:
    """
    Une circulation (date, n° de train) reconstituée à partir de TOUTES ses
    lignes CSV (OUI et NON) : gares desservies, heures, ordre de desserte.
    """
    __slots__ = ("date_str", "train_no", "dep", "arr", "edges", "_reach")

    def __init__(self, date_str: str, train_no: str) -> None:
        self.date_str = date_str
        self.train_no = train_no
        self.dep: Dict[str, datetime] = {}      # heure de départ à la gare
        self.arr: Dict[str, datetime] = {}      # heure d'arrivée à la gare
        self.edges: Dict[str, Set[str]] = defaultdict(set)
        self._reach: Dict[str, Set[str]] = {}

    def add_row(self, o: str, d: str, dt_o: datetime, dt_d: datetime) -> None:
        self.edges[o].add(d)
        self.dep.setdefault(o, dt_o)
        self.arr.setdefault(d, dt_d)
        self._reach.clear()

    @property
    def stations(self) -> Set[str]:
        return set(self.dep) | set(self.arr)

    def reachable(self, s: str) -> Set[str]:
        """Gares desservies après `s` (sur la même branche)."""
        if s not in self._reach:
            seen: Set[str] = set()
            stack = list(self.edges.get(s, ()))
            while stack:
                n = stack.pop()
                if n not in seen:
                    seen.add(n)
                    stack.extend(self.edges.get(n, ()))
            self._reach[s] = seen
        return self._reach[s]

    def time_key(self, s: str) -> datetime:
        return self.dep.get(s) or self.arr[s]

    def stops_between(self, o: str, d: str) -> List[str]:
        """
        Gares desservies de o à d inclus, dans l'ordre.
        Garde-fou branches : d doit être atteignable depuis s (une gare de
        l'autre branche n'a jamais de ligne vers d). Et s doit suivre o :
        atteignable depuis o, ou à défaut située entre o et d dans le temps
        (cas d'une gare qui n'apparaît qu'en origine dans le CSV).
        """
        r_o = self.reachable(o)
        t_o, t_d = self.time_key(o), self.time_key(d)
        mids = [s for s in self.stations
                if s != d and s != o and d in self.reachable(s)
                and (s in r_o or t_o < self.time_key(s) < t_d)]
        mids.sort(key=lambda s: (self.time_key(s), s))
        return [o] + mids + [d]

    def boarding_time(self, s: str) -> Tuple[datetime, bool]:
        """(heure, exacte ?) — départ si connu, sinon arrivée (arrêt de ~2 min)."""
        if s in self.dep:
            return self.dep[s], True
        return self.arr[s], False

    def alighting_time(self, s: str) -> Tuple[datetime, bool]:
        if s in self.arr:
            return self.arr[s], True
        return self.dep[s], False


@dataclass
class Dataset:
    tickets: List[MaxTrip] = field(default_factory=list)
    tickets_by_date: Dict[str, List[MaxTrip]] = field(default_factory=dict)
    runs: Dict[Tuple[str, str], TrainRun] = field(default_factory=dict)
    stations: Set[str] = field(default_factory=set)
    n_rows: int = 0
    debug: dict = field(default_factory=dict)

    @property
    def dates(self) -> List[str]:
        return sorted(self.tickets_by_date)


def _find_column(headers: Sequence[str], keywords: Sequence[str],
                 exclude: Sequence[str] = ()) -> Optional[int]:
    """Colonne contenant tous les mots-clés et aucun exclu (la plus courte)."""
    low = [normalize(h) for h in headers]
    keys = [normalize(k) for k in keywords]
    excl = [normalize(e) for e in exclude]
    cands = sorted((len(h), i) for i, h in enumerate(low)
                   if all(k in h for k in keys) and not any(e in h for e in excl))
    return cands[0][1] if cands else None


def _parse_hhmm(s: str) -> Tuple[int, int]:
    m = re.match(r"^\s*(\d{1,2}):(\d{2})", s or "")
    if not m:
        raise ValueError(f"Heure invalide : {s!r}")
    return int(m.group(1)), int(m.group(2))


def is_cache_fresh(path: Path = TGVMAX_CACHE_PATH,
                   max_age_hours: int = CACHE_MAX_AGE_HOURS) -> bool:
    return path.exists() and (time.time() - path.stat().st_mtime) / 3600 < max_age_hours


def download_tgvmax(dest: Path = TGVMAX_CACHE_PATH, progress_cb=None) -> None:
    req = urllib.request.Request(TGVMAX_URL, headers={"User-Agent": "max-jeune-app/2.0"})
    tmp = dest.with_suffix(dest.suffix + ".tmp")
    total = 0
    with urllib.request.urlopen(req, timeout=180) as resp, open(tmp, "wb") as out:
        while True:
            chunk = resp.read(1 << 16)
            if not chunk:
                break
            out.write(chunk)
            total += len(chunk)
            if progress_cb:
                progress_cb(total)
    os.replace(tmp, dest)


class CsvRow(NamedTuple):
    """Une ligne du CSV TGVMax, parsée et normalisée."""
    date_str: str
    train: str
    axe: str
    o: str
    d: str
    dt_d: datetime
    dt_a: datetime
    oui: bool


def iter_csv_rows(path: Path = TGVMAX_CACHE_PATH, debug: Optional[dict] = None) -> Iterable[CsvRow]:
    """
    Lit le CSV TGVMax ligne à ligne (OUI + NON), dans l'ordre du fichier.
    Utilisé par load_dataset (appli) et par le pipeline web (build_data.py).
    """
    if not path.exists():
        raise FileNotFoundError(f"Cache TGVMax absent ({path}). Clique 'Recharger données'.")
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f, delimiter=";")
        headers = next(reader, None)
        if not headers:
            return
        col = {
            "date": _find_column(headers, ["DATE"]),
            "train": _find_column(headers, ["TRAIN"]),
            "axe": _find_column(headers, ["AXE"]),
            "o": _find_column(headers, ["ORIGINE"], exclude=["IATA"]),
            "d": _find_column(headers, ["DESTINATION"], exclude=["IATA"]),
            "o_iata": _find_column(headers, ["ORIGINE", "IATA"]),
            "d_iata": _find_column(headers, ["DESTINATION", "IATA"]),
            "hdep": _find_column(headers, ["HEURE", "DEPART"]),
            "harr": _find_column(headers, ["HEURE", "ARRIVEE"]),
            "dispo": _find_column(headers, ["DISPONIBILITE"]) or _find_column(headers, ["JEUNE"]),
        }
        required = ["date", "train", "o", "d", "hdep", "harr", "dispo"]
        missing = [k for k in required if col[k] is None]
        if missing:
            raise RuntimeError(f"Colonnes manquantes : {missing}\nEntêtes : {headers}")
        if debug is not None:
            debug["headers"] = list(headers)
            debug["columns"] = {k: (v, headers[v] if v is not None else None) for k, v in col.items()}
        width = max(v for v in col.values() if v is not None)

        def get(row, k):
            i = col[k]
            return row[i].strip() if i is not None and i < len(row) else ""

        day_cache: Dict[str, datetime] = {}
        for row in reader:
            if len(row) <= width:
                continue
            date_str = get(row, "date")
            try:
                day = day_cache.get(date_str)
                if day is None:
                    day = day_cache[date_str] = datetime.strptime(date_str, "%Y-%m-%d")
                hd, md = _parse_hhmm(get(row, "hdep"))
                ha, ma = _parse_hhmm(get(row, "harr"))
            except ValueError:
                continue
            dt_d = day + timedelta(hours=hd, minutes=md)
            dt_a = day + timedelta(hours=ha, minutes=ma)
            if dt_a < dt_d:                      # passage de minuit
                dt_a += timedelta(days=1)
            yield CsvRow(date_str, get(row, "train"), get(row, "axe"),
                         resolve_csv_station(get(row, "o"), get(row, "o_iata")),
                         resolve_csv_station(get(row, "d"), get(row, "d_iata")),
                         dt_d, dt_a, get(row, "dispo").upper() == "OUI")


def load_dataset(path: Path = TGVMAX_CACHE_PATH) -> Dataset:
    """
    Lit TOUT le CSV (OUI + NON).
      - lignes OUI  -> billets MAX réservables (Dataset.tickets)
      - toutes      -> circulations reconstituées (Dataset.runs)
    """
    ds = Dataset()
    n_oui = 0
    for r in iter_csv_rows(path, ds.debug):
        ds.n_rows += 1
        key = (r.date_str, r.train)
        run = ds.runs.get(key)
        if run is None:
            run = ds.runs[key] = TrainRun(r.date_str, r.train)
        run.add_row(r.o, r.d, r.dt_d, r.dt_a)
        ds.stations.add(r.o)
        ds.stations.add(r.d)
        if r.oui:
            n_oui += 1
            t = MaxTrip(r.date_str, r.train, r.axe, r.o, r.d,
                        f"{r.dt_d:%H:%M}", f"{r.dt_a:%H:%M}", r.dt_d, r.dt_a)
            ds.tickets.append(t)
            ds.tickets_by_date.setdefault(r.date_str, []).append(t)
    ds.debug["n_oui"] = n_oui
    return ds


# =============================================================================
# 6. RÉSOLUTION DES SAISIES UTILISATEUR
# =============================================================================

HUB_SUFFIX = " (HUB)"
HUB_EXT_SUFFIX = " (HUB ETENDU)"


def build_autocomplete_entries(stations: Iterable[str]) -> List[str]:
    items = set(stations)
    for name in HUBS:
        if name.endswith(" ETENDU"):
            items.add(name[: -len(" ETENDU")] + HUB_EXT_SUFFIX)
        else:
            items.add(name + HUB_SUFFIX)
    return sorted(items)


@dataclass
class ResolvedLocation:
    display: str
    candidates: Set[str]
    is_hub: bool = False


def _hub_location(base: str, extended: bool) -> ResolvedLocation:
    if extended and base in HUB_EXTENDED_MAP:
        return ResolvedLocation(base + HUB_EXT_SUFFIX, set(HUBS[HUB_EXTENDED_MAP[base]]), True)
    return ResolvedLocation(base + HUB_SUFFIX, set(HUBS[base]), True)


def resolve_location(user_text: str, known_stations: Set[str],
                     include_extended: bool = False) -> Optional[ResolvedLocation]:
    """
    Saisie -> gares candidates.
      "LYON (HUB ETENDU)" -> hub étendu ; "LYON (HUB)" ou "LYON" -> hub
      (étendu si l'option est cochée) ; sinon gare exacte, alias, ou
      recherche approximative.
    """
    if not user_text or not user_text.strip():
        return None
    raw = user_text.strip()
    up = raw.upper()
    if up.endswith(HUB_EXT_SUFFIX):
        base = normalize(raw[: -len(HUB_EXT_SUFFIX)])
        if base in HUB_EXTENDED_MAP:
            return _hub_location(base, True)
    if up.endswith(HUB_SUFFIX):
        base = normalize(raw[: -len(HUB_SUFFIX)])
        if base in HUBS:
            return _hub_location(base, include_extended)
    n = normalize(raw)
    if n in HUBS:                       # avant les alias (ex. "MARSEILLE")
        return _hub_location(n, include_extended)
    n = resolve_alias(raw)
    if n in known_stations:
        return ResolvedLocation(n, {n})
    cands = [s for s in known_stations if s.startswith(n)] or [s for s in known_stations if n in s]
    if len(cands) == 1:
        return ResolvedLocation(cands[0], {cands[0]})
    if cands:
        return ResolvedLocation(f"{n} (plusieurs gares)", set(cands))
    return None


# =============================================================================
# 7. TRONÇONS RÉSERVABLES ("legs")
# =============================================================================
#
# Un leg = un billet MAX OUI + la gare où l'on monte + la gare où l'on
# descend. Si montée = origine officielle et descente = destination
# officielle, c'est un MAX DIRECT ; sinon une montée/descente en cours.

@dataclass(frozen=True)
class Leg:
    ticket: MaxTrip
    board: str
    alight: str
    dt_board: datetime
    dt_alight: datetime
    exact_board: bool = True
    exact_alight: bool = True

    @property
    def is_direct(self) -> bool:
        return self.board == self.ticket.origine and self.alight == self.ticket.destination

    @property
    def ticket_minutes(self) -> float:
        t = self.ticket
        return (t.datetime_arrivee - t.datetime_depart).total_seconds() / 60


def _ticket_stops(ds: Dataset, t: MaxTrip) -> Tuple[Optional[TrainRun], List[str]]:
    run = ds.runs.get((t.date_str, t.train_no))
    if run is None:
        return None, [t.origine, t.destination]
    return run, run.stops_between(t.origine, t.destination)


def _make_leg(t: MaxTrip, run: Optional[TrainRun], b: str, a: str) -> Leg:
    if b == t.origine or run is None:
        dtb, eb = t.datetime_depart, True
    else:
        dtb, eb = run.boarding_time(b)
    if a == t.destination or run is None:
        dta, ea = t.datetime_arrivee, True
    else:
        dta, ea = run.alighting_time(a)
    return Leg(t, b, a, dtb, dta, eb, ea)


def find_legs(ds: Dataset, date_str: str,
              dep_cands: Optional[Set[str]], arr_cands: Optional[Set[str]]) -> List[Leg]:
    """
    Tous les tronçons réservables du jour allant d'une gare de `dep_cands`
    à une gare de `arr_cands` (None = n'importe quelle gare).
    Quand arr_cands est fourni, on ne garde que la 1re gare d'arrivée
    atteinte (arrivée la plus tôt) pour chaque gare de montée.
    """
    legs: List[Leg] = []
    for t in ds.tickets_by_date.get(date_str, ()):
        if dep_cands is not None and arr_cands is not None \
                and t.origine in dep_cands and t.destination in arr_cands:
            legs.append(_make_leg(t, None, t.origine, t.destination))
            continue
        run, stops = _ticket_stops(ds, t)
        for i, b in enumerate(stops[:-1]):
            if dep_cands is not None and b not in dep_cands:
                continue
            for a in stops[i + 1:]:
                if arr_cands is not None and a not in arr_cands:
                    continue
                legs.append(_make_leg(t, run, b, a))
                if arr_cands is not None:
                    break
    return legs


# =============================================================================
# 8. RECHERCHE : Parties A, B, C
# =============================================================================

GROUP_A, GROUP_B, GROUP_C, GROUP_D = "A", "B", "C", "D"
GROUP_LABELS = {
    GROUP_A: "MAX DIRECT",
    GROUP_B: "MONTÉE / DESCENTE EN COURS",
    GROUP_C: "2 SEGMENTS MAX",
    GROUP_D: "MAX + TER",
}


@dataclass
class Query:
    date_str: str
    dep_cands: Set[str]
    arr_cands: Set[str]
    heure_min: str = "00:00"
    arrivee_max: Optional[str] = None      # "HH:MM" ou None
    allow_boarding: bool = True
    allow_two_segments: bool = True
    allow_ter: bool = True          # Partie D, si aucune solution 100 % MAX
    force_ter: bool = False         # Partie D affichée même s'il existe du 100 % MAX


@dataclass
class SearchResult:
    group: str
    itineraire: str
    dt_dep: datetime
    dt_arr: datetime
    train_no: str
    od_max: str
    commentaire: str
    date_str: str
    exact_dep: bool = True
    exact_arr: bool = True

    @property
    def type(self) -> str:
        return GROUP_LABELS[self.group]

    @property
    def depart(self) -> str:
        return ("" if self.exact_dep else "~") + self.dt_dep.strftime("%H:%M")

    @property
    def arrivee(self) -> str:
        s = self.dt_arr.strftime("%H:%M")
        if self.dt_arr.date() > self.dt_dep.date():
            s += " (+1j)"
        return ("" if self.exact_arr else "~") + s

    @property
    def minutes(self) -> int:
        return int((self.dt_arr - self.dt_dep).total_seconds() // 60)

    @property
    def duree(self) -> str:
        m = self.minutes
        return f"{m // 60}h{m % 60:02d}"


def result_sort_key(r: SearchResult):
    """Ordre d'affichage par défaut, totalement déterministe."""
    return (r.dt_dep, r.dt_arr, r.train_no, r.itineraire)


@dataclass
class SearchOutcome:
    direct: List[SearchResult]
    boarding: List[SearchResult]
    two_segments: List[SearchResult]
    max_ter: List[SearchResult] = field(default_factory=list)

    @property
    def all_max(self) -> List[SearchResult]:
        return self.direct + self.boarding + self.two_segments

    @property
    def all(self) -> List[SearchResult]:
        return self.all_max + self.max_ter


def _minutes_of(hhmm: Optional[str]) -> Optional[int]:
    if not hhmm:
        return None
    try:
        h, m = _parse_hhmm(hhmm)
        return h * 60 + m
    except ValueError:
        return None


def _time_ok(dt_board: datetime, dt_alight: datetime, q: Query) -> bool:
    day = datetime.strptime(q.date_str, "%Y-%m-%d")
    hmin = _minutes_of(q.heure_min) or 0
    if dt_board < day + timedelta(minutes=hmin):
        return False
    amax = _minutes_of(q.arrivee_max)
    if amax is not None and dt_alight > day + timedelta(minutes=amax):
        return False
    return True


def _best_per_ride(legs: Iterable[Leg]) -> Dict[Tuple[str, str, str], Tuple[Leg, List[Leg]]]:
    """
    Regroupe les legs d'un même trajet physique (train, montée, descente).
    Plusieurs billets OUI peuvent le couvrir (ex. Paris→Toulon et
    Paris→St-Raphaël sur le même train) : on garde le billet à l'OD la plus
    courte et on conserve les autres pour information.
    """
    groups: Dict[Tuple[str, str, str], List[Leg]] = defaultdict(list)
    for l in legs:
        groups[(l.ticket.train_no, l.board, l.alight)].append(l)
    out = {}
    for k, ls in groups.items():
        ls.sort(key=lambda l: (not l.is_direct, l.ticket_minutes, l.ticket.origine,
                               l.ticket.destination))
        out[k] = (ls[0], ls[1:])
    return out


def _cap(s: str) -> str:
    """Majuscule initiale sans toucher au reste (str.capitalize minusculerait 'PARIS')."""
    return s[:1].upper() + s[1:]


def _od(t: MaxTrip) -> str:
    return f"{t.origine} → {t.destination}"


def search_direct_and_boarding(ds: Dataset, q: Query) -> Tuple[List[SearchResult], List[SearchResult]]:
    """Parties A (MAX direct) et B (montée / descente en cours d'arrêt)."""
    legs = [l for l in find_legs(ds, q.date_str, q.dep_cands, q.arr_cands)
            if _time_ok(l.dt_board, l.dt_alight, q)]
    rides = _best_per_ride(legs)

    direct: List[SearchResult] = []
    by_train_b: Dict[str, List[Tuple[Leg, List[Leg]]]] = defaultdict(list)
    for (train, _, _), (leg, others) in rides.items():
        if leg.is_direct:
            # Même train, mêmes gares : les billets plus longs sont signalés
            # mais pas proposés en doublon.
            comment = "Billet MAX direct."
            if others:
                comment += (" Aussi couvert par les billets MAX : "
                            + ", ".join(_od(o.ticket) for o in others) + ".")
            direct.append(SearchResult(
                GROUP_A, f"{leg.board} → {leg.alight}", leg.dt_board, leg.dt_alight,
                leg.ticket.train_no, _od(leg.ticket), comment, q.date_str))
        else:
            # Le masquage ne porte que sur la même paire de gares (géré par
            # _best_per_ride) : monter à une autre gare du hub reste proposé.
            by_train_b[train].append((leg, others))

    boarding: List[SearchResult] = []
    if q.allow_boarding:
        for train, options in by_train_b.items():
            # Un seul résultat par train : arrivée la plus tôt, puis montée la
            # plus tardive ; les autres gares de montée sont signalées.
            options.sort(key=lambda o: (o[0].dt_alight, -o[0].dt_board.timestamp(),
                                        o[0].board, o[0].alight))
            leg, others = options[0]
            t = leg.ticket
            up, down = leg.board != t.origine, leg.alight != t.destination
            case = ("monter ET descendre en cours de route" if up and down
                    else "monter en cours de route" if up
                    else "descendre avant le terminus")
            parts = [f"{_cap(case)} : réserver le billet MAX {_od(t)} "
                     f"({t.heure_depart}→{t.heure_arrivee})."]
            if not (leg.exact_board and leg.exact_alight):
                parts.append("Heure ~ : heure d'arrivée en gare (arrêt de quelques minutes).")
            if others:
                parts.append("Autres billets valables : " + ", ".join(_od(o.ticket) for o in others) + ".")
            alt = [o[0] for o in options[1:]]
            if alt:
                parts.append("Autre(s) montée(s) possible(s) : " + ", ".join(
                    f"{a.board} {a.dt_board:%H:%M}" for a in alt) + ".")
            boarding.append(SearchResult(
                GROUP_B, f"{leg.board} → {leg.alight}", leg.dt_board, leg.dt_alight,
                t.train_no, _od(t), " ".join(parts), q.date_str,
                leg.exact_board, leg.exact_alight))

    direct.sort(key=result_sort_key)
    boarding.sort(key=result_sort_key)
    return direct, boarding


def _connection_kind(l1: Leg, l2: Leg, margin: float) -> Optional[str]:
    """Applique les règles de marge. Renvoie le libellé ou None si invalide."""
    if margin < 0 or margin > MARGIN_CONNECTION_MAX:
        return None
    same_station = l1.alight == l2.board
    if same_station and l1.ticket.train_no == l2.ticket.train_no:
        return "même train (2 billets)" if margin <= MARGIN_SAME_TRAIN_MAX else None
    if same_station:
        return "correspondance même gare" if margin >= MARGIN_SAME_STATION_MIN else None
    if margin >= MARGIN_HUB_MIN:
        return f"changement de gare (hub {STATION_TO_HUB.get(l1.alight, '?')})"
    return None


def search_two_segments(ds: Dataset, q: Query,
                        better_options: Sequence[SearchResult] = ()) -> List[SearchResult]:
    """
    Partie C : deux billets MAX successifs. Chaque segment peut lui-même être
    une montée/descente en cours d'arrêt (arrêts réels issus du CSV).
    Seules les solutions non dominées sont renvoyées : on écarte une
    solution s'il en existe une autre partant au plus tôt au même moment et
    arrivant au plus tard au même moment (y compris un direct / une montée).
    """
    day = datetime.strptime(q.date_str, "%Y-%m-%d")
    hmin = day + timedelta(minutes=_minutes_of(q.heure_min) or 0)

    first = [l for l in find_legs(ds, q.date_str, q.dep_cands, None)
             if l.dt_board >= hmin and l.alight not in q.arr_cands]
    second = [l for l in find_legs(ds, q.date_str, None, q.arr_cands)
              if l.board not in q.dep_cands]
    first_rides = [v[0] for v in _best_per_ride(first).values()]
    by_board: Dict[str, List[Leg]] = defaultdict(list)
    for l, _ in _best_per_ride(second).values():
        by_board[l.board].append(l)

    raw: List[Tuple[Leg, Leg, str, float]] = []
    for l1 in first_rides:
        for pivot in sorted(HUB_SIBLINGS.get(l1.alight, set()) | {l1.alight}):
            for l2 in by_board.get(pivot, ()):
                if l2.ticket is l1.ticket or l2.alight == l1.board:
                    continue
                margin = (l2.dt_board - l1.dt_alight).total_seconds() / 60
                kind = _connection_kind(l1, l2, margin)
                if kind is None:
                    continue
                total = (l2.dt_alight - l1.dt_board).total_seconds() / 60
                if total > TWO_SEG_TOTAL_DURATION_MAX_MIN:
                    continue
                if not _time_ok(l1.dt_board, l2.dt_alight, q):
                    continue
                raw.append((l1, l2, kind, margin))

    # Dominance : trié par départ décroissant, on garde une solution si son
    # arrivée est strictement meilleure que tout ce qui part au moins aussi tard.
    ref = [(r.dt_dep, r.dt_arr) for r in better_options]
    raw.sort(key=lambda x: (-x[0].dt_board.timestamp(), x[1].dt_alight,
                            "hub" in x[2], -x[3], x[0].alight, x[1].board,
                            x[0].ticket.train_no, x[1].ticket.train_no, x[0].board, x[1].alight))
    kept: List[Tuple[Leg, Leg, str, float]] = []
    best_arr: Optional[datetime] = None
    for item in raw:
        l1, l2 = item[0], item[1]
        if best_arr is not None and l2.dt_alight >= best_arr:
            continue
        if any(d >= l1.dt_board and a <= l2.dt_alight for d, a in ref):
            continue
        kept.append(item)
        best_arr = l2.dt_alight

    results: List[SearchResult] = []
    for l1, l2, kind, margin in kept:
        via = l1.alight if l1.alight == l2.board else f"{l1.alight} → {l2.board}"

        def seg(l: Leg) -> str:
            s = f"{l.board} {l.dt_board:%H:%M} → {l.alight} {l.dt_alight:%H:%M} (train {l.ticket.train_no}"
            if not l.is_direct:
                s += f", billet {_od(l.ticket)}"
            return s + ")"

        results.append(SearchResult(
            GROUP_C, f"{l1.board} → {l2.alight} (via {via})",
            l1.dt_board, l2.dt_alight,
            f"{l1.ticket.train_no}+{l2.ticket.train_no}",
            f"{_od(l1.ticket)} | {_od(l2.ticket)}",
            f"{_cap(kind)}, {int(round(margin))} min. Seg1 : {seg(l1)}. Seg2 : {seg(l2)}.",
            q.date_str, l1.exact_board, l2.exact_alight))
    results.sort(key=result_sort_key)
    return results


# =============================================================================
# 9. PARTIE D — MAX + TER (horaires théoriques du GTFS SNCF)
# =============================================================================
#
# Source : GTFS officiel SNCF (TER, Intercités, TGV), mis à jour chaque jour.
#   https://transport.data.gouv.fr/datasets/horaires-sncf  — licence ODbL
# On ne garde que les circulations non-MAX acceptées en complément :
#   Train TER, TramTrain, Car TER, Intercités, Intercités de nuit.
#
# Principe :
#   - après : billet MAX (A ou B) de la gare de départ vers une gare X, puis
#     TER direct de X (ou d'une gare sœur du hub) vers la gare d'arrivée ;
#   - avant : TER direct de la gare de départ vers une gare Y, puis billet
#     MAX de Y vers la gare d'arrivée.
#   - un seul TER (pas de changement TER-TER) ;
#   - le segment MAX ne doit pas déjà arriver (ou partir) dans la ville /
#     le hub demandé : sinon ce n'est pas un "MAX + TER" ;
#   - marges : 10 min même gare, 60 min changement de gare dans un hub ;
#   - le TER est une proposition d'horaire théorique (billet TER payant) ;
#     le segment MAX reste toujours adossé à une ligne OUI du CSV.
#   - déclenchement : seulement si aucune solution 100 % MAX n'existe dans la
#     plage horaire, sauf si force_ter est coché.

GTFS_URL = "https://eu.ftp.opendatasoft.com/sncf/plandata/Export_OpenData_SNCF_GTFS_NewTripId.zip"
GTFS_PATH = BASE_DIR / "Export_OpenData_SNCF_GTFS_NewTripId.zip"

# Modes (préfixe des StopPoint "StopPoint:OCE<mode>-<UIC>") -> libellé affiché
TER_MODES: Dict[str, str] = {
    "Train TER": "TER",
    "TramTrain": "Tram-train",
    "Car TER": "Car TER",
    "INTERCITES": "Intercités",
    "INTERCITES de nuit": "Intercités de nuit",
}
TER_SEARCH_WINDOW_MIN = 240      # on cherche un TER dans les 4 h qui suivent / précèdent

# Gares dont le libellé GTFS diffère trop du libellé TGVMax (clé = code UIC).
GTFS_UIC_TO_CANONICAL: Dict[str, str] = {
    "87686006": "PARIS GARE DE LYON",
    "87391003": "PARIS MONTPARNASSE",
    "87271007": "PARIS NORD",
    "87113001": "PARIS EST",
    "87547000": "PARIS AUSTERLITZ",
    "87713040": "DIJON VILLE",
    "87725689": "MACON VILLE",
    "87192039": "METZ VILLE",
    "87182063": "MULHOUSE VILLE",
    "87485003": "LA ROCHELLE VILLE",
    "87317586": "BOULOGNE VILLE",
    "87713412": "DOLE VILLE",
    "87764001": "MONTELIMAR GARE SNCF",
    "87745000": "BELLEGARDE SUR VALSERINE GARE",
    "87144014": "ST DIE",
    "87485490": "ST MAIXENT",
    "87784264": "PORT VENDRES VILLE",
    "87543017": "LES AUBRAIS ORLEANS",
    "87175240": "LEROUVILLE CENTRE",
}


def _match_key(name: str) -> str:
    """Clé de rapprochement de libellés (SAINT/ST, tirets, casse, accents)."""
    n = canonicalize_station(name).replace("-", " ")
    n = re.sub(r"\bSAINTE\b", "STE", n)
    n = re.sub(r"\bSAINT\b", "ST", n)
    return re.sub(r"\s+", " ", n).strip()


def _gtfs_minutes(s: str) -> Optional[int]:
    """'25:10:00' -> 1510 (les heures GTFS peuvent dépasser 24 h)."""
    m = re.match(r"^\s*(\d+):(\d{2})", s or "")
    return int(m.group(1)) * 60 + int(m.group(2)) if m else None


@dataclass(frozen=True)
class TerRide:
    """Un trajet TER/Intercités direct, horaires théoriques."""
    mode: str
    train_no: str
    board: str
    alight: str
    dt_dep: datetime
    dt_arr: datetime

    def label(self) -> str:
        return (f"{self.mode} {self.train_no} {self.board} {self.dt_dep:%H:%M} → "
                f"{self.alight} {self.dt_arr:%H:%M}")


class TerNetwork:
    """
    Réseau TER/Intercités chargé depuis le GTFS, avec gares renommées en noms
    canoniques (identiques au CSV TGVMax quand la gare y existe).
    Les index par jour sont construits à la demande et mis en cache.
    """

    def __init__(self) -> None:
        # trip = (mode, train_no, service_id, [(gare, arr_min, dep_min, montée_ok, descente_ok)])
        self.trips: List[Tuple[str, str, str, List[Tuple[str, int, int, bool, bool]]]] = []
        self.service_dates: Dict[str, Set[str]] = defaultdict(set)   # service -> {YYYY-MM-DD}
        self.stations: Set[str] = set()
        self.feed_version: str = ""
        self._dep_index: Dict[str, Dict[str, List[Tuple[datetime, int, int]]]] = {}
        self._arr_index: Dict[str, Dict[str, List[Tuple[datetime, int, int]]]] = {}

    # --- index par jour ---------------------------------------------------
    def _build_day(self, date_str: str) -> None:
        day = datetime.strptime(date_str, "%Y-%m-%d")
        dep: Dict[str, List[Tuple[datetime, int, int]]] = defaultdict(list)
        arr: Dict[str, List[Tuple[datetime, int, int]]] = defaultdict(list)
        for ti, (_, _, service, stops) in enumerate(self.trips):
            if date_str not in self.service_dates.get(service, ()):
                continue
            for pos, (st, a, d, can_board, can_alight) in enumerate(stops):
                if can_board and pos < len(stops) - 1:
                    dep[st].append((day + timedelta(minutes=d), ti, pos))
                if can_alight and pos > 0:
                    arr[st].append((day + timedelta(minutes=a), ti, pos))
        for idx in (dep, arr):
            for lst in idx.values():
                lst.sort()
        self._dep_index[date_str] = dep
        self._arr_index[date_str] = arr

    def _index(self, date_str: str, which: str):
        if date_str not in self._dep_index:
            self._build_day(date_str)
        return (self._dep_index if which == "dep" else self._arr_index)[date_str]

    def _ride(self, day: datetime, ti: int, i: int, j: int) -> TerRide:
        mode, train, _, stops = self.trips[ti]
        return TerRide(mode, train, stops[i][0], stops[j][0],
                       day + timedelta(minutes=stops[i][2]), day + timedelta(minutes=stops[j][1]))

    # --- requêtes ----------------------------------------------------------
    def earliest_arrival(self, date_str: str, from_station: str, to_set: Set[str],
                         not_before: datetime) -> Optional[TerRide]:
        """TER direct partant de from_station après not_before, arrivant le plus tôt dans to_set."""
        day = datetime.strptime(date_str, "%Y-%m-%d")
        lst = self._index(date_str, "dep").get(from_station, [])
        limit = not_before + timedelta(minutes=TER_SEARCH_WINDOW_MIN)
        best: Optional[TerRide] = None
        k = bisect_left(lst, (not_before, -1, -1))
        for dt, ti, pos in lst[k:]:
            if dt > limit or (best and dt >= best.dt_arr):
                break
            stops = self.trips[ti][3]
            for j in range(pos + 1, len(stops)):
                if stops[j][0] in to_set and stops[j][4]:
                    ride = self._ride(day, ti, pos, j)
                    if best is None or ride.dt_arr < best.dt_arr:
                        best = ride
                    break
        return best

    def latest_departure(self, date_str: str, from_set: Set[str], to_station: str,
                         not_after: datetime) -> Optional[TerRide]:
        """TER direct depuis from_set arrivant à to_station avant not_after, parti le plus tard."""
        day = datetime.strptime(date_str, "%Y-%m-%d")
        lst = self._index(date_str, "arr").get(to_station, [])
        limit = not_after - timedelta(minutes=TER_SEARCH_WINDOW_MIN)
        best: Optional[TerRide] = None
        k = bisect_right(lst, (not_after, 1 << 30, 1 << 30))
        for dt, ti, pos in reversed(lst[:k]):
            if dt < limit or (best and dt <= best.dt_dep):
                break
            stops = self.trips[ti][3]
            for i in range(pos - 1, -1, -1):
                if stops[i][0] in from_set and stops[i][3]:
                    ride = self._ride(day, ti, i, pos)
                    if best is None or ride.dt_dep > best.dt_dep:
                        best = ride
                    break
        return best


def download_gtfs(dest: Path = GTFS_PATH, progress_cb=None) -> None:
    """Télécharge le zip GTFS SNCF (≈ 5 Mo)."""
    req = urllib.request.Request(GTFS_URL, headers={"User-Agent": "max-jeune-app/2.0"})
    tmp = dest.with_suffix(dest.suffix + ".tmp")
    total = 0
    with urllib.request.urlopen(req, timeout=180) as resp, open(tmp, "wb") as out:
        while True:
            chunk = resp.read(1 << 16)
            if not chunk:
                break
            out.write(chunk)
            total += len(chunk)
            if progress_cb:
                progress_cb(total)
    os.replace(tmp, dest)


def load_ter_network(path: Path = GTFS_PATH,
                     max_stations: Iterable[str] = ()) -> TerNetwork:
    """
    Lit le zip GTFS SNCF sans le décompresser. `max_stations` (gares du CSV
    TGVMax) sert à nommer les gares du GTFS exactement comme dans le CSV.
    """
    if not path.exists():
        raise FileNotFoundError(f"GTFS absent ({path.name}). Clique 'Recharger données'.")
    net = TerNetwork()
    max_by_key = {_match_key(s): s for s in max_stations}

    def rows(zf, name):
        with zf.open(name) as f:
            yield from csv.DictReader(io.TextIOWrapper(f, encoding="utf-8-sig"))

    with zipfile.ZipFile(path) as zf:
        for r in rows(zf, "feed_info.txt"):
            net.feed_version = r.get("feed_version", "")
        # Gares : StopPoint -> StopArea -> nom canonique
        area_name: Dict[str, str] = {}
        parent: Dict[str, str] = {}
        for r in rows(zf, "stops.txt"):
            if r["location_type"] == "1":
                area_name[r["stop_id"]] = r["stop_name"]
            else:
                parent[r["stop_id"]] = r["parent_station"]

        def canonical_area(area_id: str) -> str:
            uic = area_id.rsplit("OCE", 1)[-1]
            if uic in GTFS_UIC_TO_CANONICAL:
                return GTFS_UIC_TO_CANONICAL[uic]
            key = _match_key(area_name.get(area_id, area_id))
            return max_by_key.get(key, key)

        area_canon: Dict[str, str] = {}
        stop_info: Dict[str, Tuple[str, str]] = {}   # stop_id -> (mode, gare)
        for sp, area in parent.items():
            m = re.match(r"StopPoint:OCE(.+)-\d+$", sp)
            if not m or m.group(1) not in TER_MODES:
                continue
            if area not in area_canon:
                area_canon[area] = canonical_area(area)
            stop_info[sp] = (TER_MODES[m.group(1)], area_canon[area])

        trip_meta = {r["trip_id"]: (r["trip_headsign"], r["service_id"])
                     for r in rows(zf, "trips.txt")}
        for r in rows(zf, "calendar_dates.txt"):
            if r["exception_type"] == "1":
                d = r["date"]
                net.service_dates[r["service_id"]].add(f"{d[:4]}-{d[4:6]}-{d[6:]}")

        current, stops, mode = None, [], None

        def flush():
            if current and len(stops) >= 2 and current in trip_meta:
                stops.sort(key=lambda x: x[0])
                train, service = trip_meta[current]
                net.trips.append((mode, train, service, [s[1] for s in stops]))

        for r in rows(zf, "stop_times.txt"):
            if r["trip_id"] != current:
                flush()
                current, stops, mode = r["trip_id"], [], None
            info = stop_info.get(r["stop_id"])
            if info is None:
                continue
            a, d = _gtfs_minutes(r["arrival_time"]), _gtfs_minutes(r["departure_time"])
            if a is None or d is None:
                continue
            mode = info[0]
            stops.append((int(r["stop_sequence"]),
                          (info[1], a, d, r["pickup_type"] != "1", r["drop_off_type"] != "1")))
        flush()
    for _, _, _, st in net.trips:
        net.stations.update(s[0] for s in st)
    return net


def _city_zone(cands: Set[str]) -> Set[str]:
    """Gares demandées + toutes les gares du même hub de base (Paris, Lyon, Marseille)."""
    zone = set(cands)
    for name, lst in HUBS.items():
        if not name.endswith(" ETENDU") and zone & set(lst):
            zone |= set(lst)
    return zone


def _pivot_margin(a: str, b: str) -> Optional[int]:
    if a == b:
        return MARGIN_SAME_STATION_MIN
    if b in HUB_SIBLINGS.get(a, ()):
        return MARGIN_HUB_MIN
    return None


def search_max_ter(ds: Dataset, net: TerNetwork, q: Query,
                   better_options: Sequence[SearchResult] = ()) -> List[SearchResult]:
    """Partie D : MAX puis TER, ou TER puis MAX (un seul TER direct)."""
    day = datetime.strptime(q.date_str, "%Y-%m-%d")
    hmin = day + timedelta(minutes=_minutes_of(q.heure_min) or 0)
    dep_zone, arr_zone = _city_zone(q.dep_cands), _city_zone(q.arr_cands)
    combos: List[Tuple[datetime, datetime, str, str, str, str, bool, bool]] = []

    def max_txt(l: Leg) -> str:
        s = (f"MAX {l.ticket.train_no} {l.board} {l.dt_board:%H:%M} → "
             f"{l.alight} {l.dt_alight:%H:%M}")
        if not l.is_direct:
            s += f" (billet {_od(l.ticket)}, {'montée' if l.board != l.ticket.origine else 'descente'} en cours)"
        return s

    # (1) MAX puis TER
    first = [l for l in find_legs(ds, q.date_str, q.dep_cands, None)
             if l.dt_board >= hmin and l.alight not in arr_zone]
    for l1, _ in _best_per_ride(first).values():
        for pivot in sorted(HUB_SIBLINGS.get(l1.alight, set()) | {l1.alight}):
            margin = _pivot_margin(l1.alight, pivot)
            ter = net.earliest_arrival(q.date_str, pivot, q.arr_cands,
                                       l1.dt_alight + timedelta(minutes=margin))
            if ter is None:
                continue
            combos.append((l1.dt_board, ter.dt_arr, f"{l1.board} → {ter.alight} (via {pivot})",
                           f"{l1.ticket.train_no}+{ter.mode} {ter.train_no}", _od(l1.ticket),
                           f"MAX puis TER : {max_txt(l1)}, puis {ter.label()} "
                           f"(correspondance {int((ter.dt_dep - l1.dt_alight).total_seconds() // 60)} min).",
                           l1.exact_board, True))

    # (2) TER puis MAX
    second = [l for l in find_legs(ds, q.date_str, None, q.arr_cands)
              if l.board not in dep_zone]
    for l2, _ in _best_per_ride(second).values():
        for pivot in sorted(HUB_SIBLINGS.get(l2.board, set()) | {l2.board}):
            margin = _pivot_margin(pivot, l2.board)
            ter = net.latest_departure(q.date_str, q.dep_cands, pivot,
                                       l2.dt_board - timedelta(minutes=margin))
            if ter is None or ter.dt_dep < hmin:
                continue
            combos.append((ter.dt_dep, l2.dt_alight, f"{ter.board} → {l2.alight} (via {pivot})",
                           f"{ter.mode} {ter.train_no}+{l2.ticket.train_no}", _od(l2.ticket),
                           f"TER puis MAX : {ter.label()}, puis {max_txt(l2)} "
                           f"(correspondance {int((l2.dt_board - ter.dt_arr).total_seconds() // 60)} min).",
                           True, l2.exact_alight))

    combos = [c for c in combos
              if (c[1] - c[0]).total_seconds() / 60 <= TWO_SEG_TOTAL_DURATION_MAX_MIN
              and _time_ok(c[0], c[1], q)]
    # Dominance (même règle que la Partie C)
    combos.sort(key=lambda c: (-c[0].timestamp(), c[1], c[2], c[3]))
    ref = [(r.dt_dep, r.dt_arr) for r in better_options]
    kept, best_arr = [], None
    for c in combos:
        if best_arr is not None and c[1] >= best_arr:
            continue
        if any(d >= c[0] and a <= c[1] for d, a in ref):
            continue
        kept.append(c)
        best_arr = c[1]
    results = [SearchResult(GROUP_D, it, dd, da, tr, od,
                            com + " Horaire TER théorique (GTFS SNCF), billet TER payant à acheter à part.",
                            q.date_str, ed, ea)
               for dd, da, it, tr, od, com, ed, ea in kept]
    results.sort(key=result_sort_key)
    return results


# =============================================================================
# 10. POINT D'ENTRÉE DE LA RECHERCHE
# =============================================================================

def run_search(ds: Dataset, q: Query, net: Optional[TerNetwork] = None) -> SearchOutcome:
    direct, boarding = search_direct_and_boarding(ds, q)
    two: List[SearchResult] = []
    if q.allow_two_segments:
        two = search_two_segments(ds, q, better_options=direct + boarding)
    out = SearchOutcome(direct, boarding, two)
    if net is not None and q.allow_ter and (q.force_ter or not out.all_max):
        out.max_ter = search_max_ter(ds, net, q, better_options=out.all_max)
    return out


# =============================================================================
# 11. BALAYAGE MULTI-JOURS
# =============================================================================

@dataclass
class DaySummary:
    date_str: str
    n_direct: int
    n_boarding: int
    n_two: int
    first_dep: str
    last_dep: str
    fastest: str
    n_ter: int = 0

    @property
    def total(self) -> int:
        return self.n_direct + self.n_boarding + self.n_two + self.n_ter


JOURS = ["lun", "mar", "mer", "jeu", "ven", "sam", "dim"]


def scan_days(ds: Dataset, q: Query, n_days: int = SCAN_DEFAULT_DAYS,
              progress_cb=None, net: Optional[TerNetwork] = None) -> List[DaySummary]:
    """Lance la même recherche sur chaque jour à partir de q.date_str."""
    start = datetime.strptime(q.date_str, "%Y-%m-%d").date()
    days = [d for d in ds.dates if 0 <= (date.fromisoformat(d) - start).days < n_days]
    out: List[DaySummary] = []
    for i, d in enumerate(days):
        res = run_search(ds, replace(q, date_str=d), net)
        allr = res.all
        out.append(DaySummary(
            d, len(res.direct), len(res.boarding), len(res.two_segments),
            min((r.depart for r in allr), default="—", key=lambda s: s.lstrip("~")),
            max((r.depart for r in allr), default="—", key=lambda s: s.lstrip("~")),
            min(allr, key=lambda r: r.minutes).duree if allr else "—",
            len(res.max_ter)))
        if progress_cb:
            progress_cb(i + 1, len(days))
    return out


def day_label(date_str: str) -> str:
    d = date.fromisoformat(date_str)
    return f"{JOURS[d.weekday()]} {d:%d/%m}"
